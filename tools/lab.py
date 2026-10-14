"""Build the MTH 105 lab workbooks (lab/*.xlsx) and assets/lab-data.js from one source.
All data are illustrative (seeded random). Run: python3 tools/lab.py"""
import json, os, random, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
R = random.Random(105)
GREEN = PatternFill('solid', fgColor='217346'); YEL = PatternFill('solid', fgColor='FFF9C4'); GREY = PatternFill('solid', fgColor='EFEFEF')

# ---------- datasets ----------
def students(n=40):
    majors = ['Business', 'Engineering', 'Design', 'IT', 'Media']; em = ['Dubai'] * 5 + ['Sharjah'] * 3 + ['Abu Dhabi', 'Ajman']
    modes = ['Car'] * 4 + ['Metro'] * 3 + ['Bus'] * 2 + ['Walk']; sat = ['Low', 'Medium', 'Medium', 'High', 'High']
    rows = []
    for i in range(n):
        study = round(max(1, min(25, R.gauss(11, 5))), 1)
        gpa = round(max(2.0, min(4.0, 2.1 + 0.075 * study + R.gauss(0, 0.25))), 2)
        commute = int(min(90, 8 + R.expovariate(1 / 18)))
        rows.append([f'S{i+1:02d}', R.choice(['F', 'M']), R.choice(majors), R.randint(1, 4), R.choice(em), R.choice(modes), commute,
                     round(max(4.5, min(9.5, R.gauss(6.8, 0.9))), 1), study, gpa, R.choice(sat)])
    return ['Student ID', 'Gender', 'Major', 'Year', 'Emirate', 'Commute mode', 'Commute (min)', 'Sleep (h)', 'Study (h/week)', 'GPA', 'Satisfaction'], rows
SH, SROWS = students()

def messy():
    h, rows = students(30); out = []
    for i, r in enumerate(rows):
        r = list(r)
        if i in (3, 11, 19): r[4] = R.choice([' dubai', 'DUBAI', 'Dubai '])
        if i in (5, 22): r[6] = r[6] * 10          # typing error
        if i in (7, 15, 26): r[7] = None            # missing
        if i == 9: r[9] = 'N/A'
        if i in (13, 24): r[5] = r[5].lower()
        out.append(r)
    return h, out
MH, MROWS = messy()

def ev_days(n=50):
    p = [0.15, 0.35, 0.30, 0.20]; rows = []
    for d in range(n):
        u = R.random(); x = 0 if u < p[0] else 1 if u < p[0] + p[1] else 2 if u < sum(p[:3]) else 3
        rows.append([d + 1, x])
    return ['Day', 'Chargers in use at noon'], rows
EH, EROWS = ev_days()
NH, NROWS = ['Student', 'Exam score'], [[i + 1, round(min(100, max(30, R.gauss(70, 10))))] for i in range(200)]
WH, WROWS = ['Household', 'Water use (L/day)'], [[i + 1, round(R.gauss(520, 60))] for i in range(36)]
def survey(n=200):
    rows = []
    for i in range(n):
        home = 'Villa' if R.random() < .5 else 'Apartment'; rec = 'Yes' if R.random() < (.6 if home == 'Villa' else .45) else 'No'
        rows.append([i + 1, home, rec])
    return ['Household', 'Home type', 'Recycles'], rows
CH, CROWS = survey()
def temps(n=30):
    rows = []
    for d in range(n):
        t = round(R.uniform(28, 44), 1); rows.append([d + 1, t, round(18 + 1.6 * t + R.gauss(0, 3.5), 1)])
    return ['Day', 'Max temperature (°C)', 'Electricity use (MWh)'], rows
TH, TROWS = temps()

# ---------- labs ----------
LABS = [
 dict(id='lab-01', week='1', title='Getting Started with Data & Excel', subs='Intro to Excel, importing data, exploring datasets', clo='', assess='Project /Exercise',
  goal='Open a data set, find your way around it and answer simple questions with formulas.',
  fx=['SUM', 'COUNT', 'COUNTA', 'MIN', 'MAX', 'Sort', 'Filter', 'Freeze Panes'],
  steps=['Download the workbook and open the <b>Data</b> sheet. Each row is one student; each column is one variable.',
         'Click cell A2, then <b>View → Freeze Panes → Freeze Top Row</b> so the headings stay visible.',
         'Click any heading, then <b>Data → Filter</b>. Filter <i>Emirate</i> to show Dubai only, then clear the filter.',
         'Sort by <i>Commute (min)</i>, largest to smallest: <b>Data → Sort</b>.',
         'On the <b>Tasks</b> sheet, type formulas in the yellow cells, e.g. <code>=COUNTA(Data!A2:A41)</code> and <code>=MAX(Data!G2:G41)</code>.',
         'Import the CSV version: <b>Data → From Text/CSV</b>, pick <i>lab-01-students.csv</i>, then <b>Load</b>.'],
  grid=dict(f='=MAX(Data!G2:G41)', cells=[['Question', 'Answer'], ['How many students?', '40'], ['Longest commute (min)', '=MAX(...)']], sel=(2, 1)),
  data=(SH, SROWS), csv=True,
  tasks=['How many students are in the data set?', 'How many are from Dubai? (use a filter, then count)', 'Shortest commute (min)', 'Longest commute (min)', 'Total study hours of all students', 'Which column is the GPA in?']),
 dict(id='lab-02', week='2', title='Formulating Questions & Exploring Variables', subs='Identifying variables, types, and measurement', clo='CLO 1', assess='Project /Exercise, Lab Exam',
  goal='Classify each variable and write questions the data can answer.',
  fx=['COUNTIF', 'Data Validation', 'Data types'],
  steps=['On the <b>Variables</b> sheet, choose a type for each variable from the drop-down: Nominal, Ordinal, Discrete or Continuous.',
         'Count categories with COUNTIF, e.g. <code>=COUNTIF(Data!F2:F41,"Metro")</code>.',
         'Write three questions the data could answer, one for each: a single categorical variable, a single numerical variable, two variables together.',
         'For each question, name the variable(s) you need and their type.'],
  grid=dict(f='=COUNTIF(Data!F2:F41,"Metro")', cells=[['Commute mode', 'Count'], ['Car', '=COUNTIF(...)'], ['Metro', '=COUNTIF(...)']], sel=(2, 1)),
  data=(SH, SROWS), variables=True,
  tasks=['Number of students who come by Car', 'By Metro', 'By Bus', 'Who Walk', 'Your question 1', 'Your question 2', 'Your question 3']),
 dict(id='lab-03', week='3', title='Visualizing Data', subs='Creating histograms, bar charts, scatterplots', clo='CLO1', assess='Project /Exercise, Lab Exam',
  goal='Choose and build the right chart for each kind of variable.',
  fx=['Insert → Histogram', 'Insert → Bar chart', 'Insert → Scatter', 'COUNTIF'],
  steps=['<b>Bar chart</b> (categorical): build a count table of <i>Commute mode</i> with COUNTIF, select it, <b>Insert → Column chart</b>.',
         '<b>Histogram</b> (numerical): select <i>Commute (min)</i>, <b>Insert → Statistic chart → Histogram</b>. Right-click the axis → <b>Format Axis</b> → set <i>Bin width</i> to 10.',
         'Describe the histogram shape: symmetric, right- or left-skewed?',
         '<b>Scatterplot</b> (two numerical): select <i>Study (h/week)</i> and <i>GPA</i>, <b>Insert → Scatter</b>. Add axis titles.',
         'Give every chart a title and axis labels.'],
  grid=dict(f='=COUNTIF(Data!F2:F41,B2)', cells=[['', 'Mode', 'Count'], ['', 'Car', '16'], ['', 'Metro', '12']], sel=(1, 2)),
  data=(SH, SROWS),
  tasks=['Shape of the commute histogram', 'Which mode is the most common?', 'Does GPA rise with study hours? (from the scatterplot)']),
 dict(id='lab-04', week='4', title='Descriptive Statistics in Excel', subs='Measurement of location, Measurement of dispersion', clo='CLO1', assess='Project /Exercise, Lab Exam',
  goal='Compute centre and spread with formulas, then check them with the Analysis ToolPak.',
  fx=['AVERAGE', 'MEDIAN', 'MODE.SNGL', 'QUARTILE.EXC', 'STDEV.S', 'VAR.S', 'Analysis ToolPak'],
  steps=['In the yellow cells type <code>=AVERAGE(Data!G2:G41)</code>, <code>=MEDIAN(...)</code>, <code>=MODE.SNGL(...)</code>.',
         'Quartiles: <code>=QUARTILE.EXC(Data!G2:G41,1)</code> and <code>,3)</code>. IQR = Q3 − Q1.',
         'Spread: <code>=VAR.S(...)</code> (sample variance), <code>=STDEV.S(...)</code>, CV = STDEV.S ÷ AVERAGE (format as %).',
         'Turn on the ToolPak once. Windows: <b>File → Options → Add-ins → Manage Excel Add-ins → Go → Analysis ToolPak</b>. Mac: <b>Tools → Excel Add-ins</b>.',
         '<b>Data → Data Analysis → Descriptive Statistics</b>, input range G1:G41, tick <i>Labels</i> and <i>Summary statistics</i>. Compare with your formulas.'],
  grid=dict(f='=STDEV.S(Data!G2:G41)', cells=[['Measure', 'Commute (min)'], ['Mean', '=AVERAGE(...)'], ['SD', '=STDEV.S(...)']], sel=(2, 1)),
  data=(SH, SROWS),
  tasks=['Mean commute', 'Median commute', 'Mode commute', 'Q1', 'Q3', 'IQR', 'Sample variance', 'Sample SD', 'CV (%)', 'Mean or median: which is the better typical value? Why?']),
 dict(id='lab-05', week='5', title='Data Cleaning & Coding', subs='Handling missing data, recoding variables', clo='CLO1', assess='Project /Exercise, Lab Exam',
  goal='Find and fix problems before analysing: blanks, typing errors, inconsistent labels.',
  fx=['TRIM', 'PROPER', 'COUNTBLANK', 'IF', 'IFS', 'ISNUMBER', 'Conditional Formatting'],
  steps=['Count missing values: <code>=COUNTBLANK(Data!H2:H31)</code>.',
         'Fix labels in a helper column: <code>=PROPER(TRIM(E2))</code> turns " dubai" into "Dubai".',
         'Spot impossible values: <b>Home → Conditional Formatting → Highlight Cells → Greater Than</b> 120 on <i>Commute (min)</i>. Decide: fix or remove.',
         'Find text in a number column: <code>=ISNUMBER(J2)</code> → FALSE for "N/A". Replace with a blank.',
         'Recode commute into groups: <code>=IFS(G2&lt;=20,"Short",G2&lt;=45,"Medium",TRUE,"Long")</code>.',
         'Write a short log of every change you made.'],
  grid=dict(f='=PROPER(TRIM(E2))', cells=[['Emirate', 'Emirate (clean)'], [' dubai', 'Dubai'], ['DUBAI', 'Dubai']], sel=(1, 1)),
  data=(MH, MROWS),
  tasks=['Number of blank Sleep values', 'Number of label problems in Emirate', 'Rows with impossible commute times', 'Cells with text in GPA', 'Your cleaning log']),
 dict(id='lab-06', week='6 & 7', title='Probability Concepts & Discrete Probability Distributions', subs='Random events, basic probability, expected value and standard deviation', clo='CLO2', assess='Project /Exercise',
  goal='Turn counts into a probability distribution and compute E(X), Var(X) and binomial probabilities.',
  fx=['COUNTIF', 'SUMPRODUCT', 'SQRT', 'BINOM.DIST', 'RAND'],
  steps=['Count each value of X: <code>=COUNTIF(Data!B2:B51,0)</code>, then 1, 2, 3.',
         'P(x) = count ÷ 50. Check that the P(x) column adds to 1.',
         'E(X): <code>=SUMPRODUCT(A2:A5,C2:C5)</code> (x column times P(x) column).',
         'Var(X): add a column (x − E)²·P(x), then SUM it. SD = <code>=SQRT(...)</code>.',
         'Binomial: P(X = 2) with n = 8, p = 0.25 → <code>=BINOM.DIST(2,8,0.25,FALSE)</code>. Use TRUE for P(X ≤ 2).',
         'Simulate a coin: <code>=IF(RAND()&lt;0.5,"H","T")</code>, fill down 100 rows, count the heads.'],
  grid=dict(f='=SUMPRODUCT(A2:A5,C2:C5)', cells=[['x', 'Count', 'P(x)'], ['0', '7', '0.14'], ['E(X)', '', '=SUMPRODUCT(...)']], sel=(3, 2)),
  data=(EH, EROWS),
  tasks=['P(0)', 'P(1)', 'P(2)', 'P(3)', 'E(X)', 'Var(X)', 'SD', 'BINOM.DIST: P(X = 2), n = 8, p = 0.25', 'P(X ≤ 2), n = 8, p = 0.25', 'Heads in your 100 simulated tosses']),
 dict(id='lab-08', week='8', title='Normal Distribution & Z-Scores', subs='NORM.DIST, NORM.S. DIST functions, calculating and interpreting Z-scores', clo='CLO2', assess='Project /Exercise',
  goal='Compute z-scores and normal probabilities, and check them against the Z-table.',
  fx=['AVERAGE', 'STDEV.S', 'STANDARDIZE', 'NORM.DIST', 'NORM.S.DIST'],
  steps=['Find μ and σ of the 200 scores: <code>=AVERAGE(Data!B2:B201)</code>, <code>=STDEV.S(...)</code>.',
         'z-score of each student: <code>=STANDARDIZE(B2,$F$1,$F$2)</code> or <code>=(B2-$F$1)/$F$2</code>. Fill down.',
         'P(X &lt; 85) with μ = 70, σ = 10: <code>=NORM.DIST(85,70,10,TRUE)</code>.',
         'Same with z: <code>=NORM.S.DIST(1.5,TRUE)</code>. Compare with the Z-table (0.9332).',
         'P(X &gt; 85) = 1 − NORM.DIST(...). P(60 &lt; X &lt; 80) = NORM.DIST(80,…) − NORM.DIST(60,…).'],
  grid=dict(f='=NORM.DIST(85,70,10,TRUE)', cells=[['Question', 'Answer'], ['P(X < 85)', '0.9332'], ['P(X > 85)', '0.0668']], sel=(1, 1)),
  data=(NH, NROWS),
  tasks=['Mean of the scores', 'SD of the scores', 'z-score of a score of 85 (using your mean and SD)', 'P(X < 85), μ = 70, σ = 10', 'P(X > 85)', 'P(60 < X < 80)', 'How many students have z above 2?']),
 dict(id='lab-09', week='9', title='Central Limit Theorem in Excel', subs='Simulating sampling distributions, using charts to visualize results', clo='CLO2', assess='',
  goal='See the CLT happen: means of skewed data become bell-shaped.',
  fx=['RAND', 'LN', 'AVERAGE', 'Histogram'],
  steps=['In <b>Sim</b>, cell A2: <code>=-10*LN(1-RAND())</code> gives a right-skewed waiting time (mean 10).',
         'Fill A2 across to AD2 (30 values = one sample of n = 30), then down to row 501 (500 samples).',
         'In AE2: <code>=AVERAGE(A2:AD2)</code>, fill down. These are 500 sample means.',
         'Make a histogram of column A, then of column AE. Compare the shapes.',
         'Compare the SD of the means, <code>=STDEV.S(AE2:AE501)</code>, with σ ÷ √n = 10 ÷ √30 ≈ 1.83.',
         'Press F9 to draw new samples and watch the histograms change.'],
  grid=dict(f='=-10*LN(1-RAND())', cells=[['x1', 'x2', '…', 'Mean'], ['3.2', '17.9', '…', '=AVERAGE(A2:AD2)'], ['8.1', '0.6', '…', '9.7']], sel=(1, 0)),
  data=None, sim=True,
  tasks=['Shape of the histogram of column A', 'Shape of the histogram of the means', 'Mean of the 500 means', 'SD of the 500 means', 'σ ÷ √n']),
 dict(id='lab-10', week='10 & 11', title='Confidence Intervals', subs='CONFIDENCE.NORM, CONFIDENCE.T functions · Manual calculation using formulas', clo='CLO3', assess='Project /Exercise, Lab Exam',
  goal='Build confidence intervals with Excel functions and by hand, and check they match.',
  fx=['AVERAGE', 'STDEV.S', 'COUNT', 'CONFIDENCE.NORM', 'CONFIDENCE.T', 'SQRT'],
  steps=['x̄ = <code>=AVERAGE(Data!B2:B37)</code>, n = <code>=COUNT(...)</code>.',
         'σ is known (60 L): margin E = <code>=CONFIDENCE.NORM(0.05,60,36)</code> for 95%.',
         'Interval: x̄ − E and x̄ + E.',
         'By hand: <code>=1.96*60/SQRT(36)</code>. It must match CONFIDENCE.NORM.',
         'σ unknown: use s = <code>=STDEV.S(...)</code> and <code>=CONFIDENCE.T(0.05,s,36)</code>. Compare the widths.',
         'Repeat with α = 0.10 and 0.01 (90% and 99%).'],
  grid=dict(f='=CONFIDENCE.NORM(0.05,60,36)', cells=[['', 'Value'], ['Mean', '=AVERAGE(...)'], ['E (95%)', '19.6']], sel=(2, 1)),
  data=(WH, WROWS),
  tasks=['x̄', 'n', 'E with CONFIDENCE.NORM (σ = 60, 95%)', 'Lower limit', 'Upper limit', 'E by hand (1.96 × 60 ÷ √36)', 's', 'E with CONFIDENCE.T (95%)', '90% interval', '99% interval']),
 dict(id='lab-12', week='12', title='Chi-Square Test in Excel', subs='Creating contingency tables · CHISQ.TEST function', clo='CLO4', assess='Project /Exercise · Due: End of Week 14',
  goal='Build a contingency table from raw data and test independence.',
  fx=['COUNTIFS', 'PivotTable', 'SUM', 'CHISQ.TEST', 'CHISQ.INV.RT'],
  steps=['Count each cell: <code>=COUNTIFS(Data!B2:B201,"Villa",Data!C2:C201,"Yes")</code>. Or use <b>Insert → PivotTable</b>.',
         'Add row totals, column totals and the grand total with SUM.',
         'Expected counts: <code>=row total*column total/grand total</code> (use $ to lock the totals).',
         'p-value: <code>=CHISQ.TEST(observed range, expected range)</code>.',
         'Critical value at α = 0.05: <code>=CHISQ.INV.RT(0.05,1)</code> = 3.841.',
         'Decide: p &lt; 0.05 → reject H₀ (the variables are related). Write the conclusion in context.'],
  grid=dict(f='=COUNTIFS(Data!B2:B201,"Villa",Data!C2:C201,"Yes")', cells=[['', 'Yes', 'No'], ['Villa', '=COUNTIFS(...)', ''], ['Apartment', '', '']], sel=(1, 1)),
  data=(CH, CROWS),
  tasks=['Observed: Villa, Yes', 'Villa, No', 'Apartment, Yes', 'Apartment, No', 'Expected: Villa, Yes', 'p-value (CHISQ.TEST)', 'Critical value (α = 0.05, df = 1)', 'Decision and conclusion']),
 dict(id='lab-13', week='13', title='Correlation Analysis', subs='CORREL, PEARSON functions, Scatterplots', clo='CLO5', assess='Project /Exercise, Lab Exam',
  goal='Measure and describe the relationship between two numerical variables.',
  fx=['CORREL', 'PEARSON', 'RSQ', 'Scatter', 'Trendline'],
  steps=['Scatterplot: select both columns, <b>Insert → Scatter</b>. x = temperature, y = electricity.',
         'r: <code>=CORREL(Data!B2:B31,Data!C2:C31)</code>. <code>=PEARSON(...)</code> gives the same value.',
         'r²: <code>=RSQ(Data!C2:C31,Data!B2:B31)</code> or r squared.',
         'Describe: direction (positive/negative) and strength (weak/moderate/strong).',
         'Right-click the points → <b>Add Trendline → Linear</b>, tick <i>Display Equation</i> and <i>R-squared</i>.'],
  grid=dict(f='=CORREL(Data!B2:B31,Data!C2:C31)', cells=[['', 'Value'], ['r', '=CORREL(...)'], ['r²', '=RSQ(...)']], sel=(1, 1)),
  data=(TH, TROWS),
  tasks=['r (CORREL)', 'r (PEARSON)', 'r²', 'Direction and strength', 'What % of the variation in electricity use is explained by temperature?']),
 dict(id='lab-14', week='14-15', title='Simple Linear Regression', subs='Using Data Analysis Toolpak: Regression, Interpreting output (model coefficients, significance model, and significance of the slope)', clo='CLO5', assess='Project /Exercise, Lab Exam · Lab Exam: End of Week 15',
  goal='Fit ŷ = a + bx, read the ToolPak output and make predictions.',
  fx=['SLOPE', 'INTERCEPT', 'FORECAST.LINEAR', 'Analysis ToolPak → Regression'],
  steps=['b: <code>=SLOPE(Data!C2:C31,Data!B2:B31)</code>. a: <code>=INTERCEPT(Data!C2:C31,Data!B2:B31)</code>. Note: y range first.',
         '<b>Data → Data Analysis → Regression</b>. Input Y: C1:C31, Input X: B1:B31, tick <i>Labels</i>, output to a new sheet.',
         'Read the output: <i>Coefficients</i> (a = Intercept, b = x row), <i>R Square</i>, <i>Significance F</i> (whole model) and the <i>P-value</i> of the slope.',
         'If the slope’s p-value &lt; 0.05, temperature is a significant predictor of electricity use.',
         'Predict at 40 °C: <code>=FORECAST.LINEAR(40,Data!C2:C31,Data!B2:B31)</code>. Check: a + b × 40.',
         'Interpret b in one sentence, in context.'],
  grid=dict(f='=SLOPE(Data!C2:C31,Data!B2:B31)', cells=[['', 'Coefficients', 'P-value'], ['Intercept', '…', '…'], ['Max temperature', '…', '…']], sel=(2, 1)),
  data=(TH, TROWS),
  tasks=['b (SLOPE)', 'a (INTERCEPT)', 'Regression equation', 'R Square (ToolPak)', 'Significance F', 'P-value of the slope', 'Is the slope significant at α = 0.05?', 'Prediction at 40 °C', 'Interpretation of b']),
]

def autosize(ws):
    for col in ws.columns:
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(60, max(10, w + 2))

def strip(html):
    import re
    return re.sub(r'<[^>]+>', '', html).replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&')

os.makedirs(os.path.join(ROOT, 'lab'), exist_ok=True)
for L in LABS:
    wb = Workbook(); ws = wb.active; ws.title = 'Instructions'
    ws['A1'] = f"MTH 105 · Lab Week {L['week']} · {L['title']}"; ws['A1'].font = Font(bold=True, size=14, color='FFFFFF'); ws['A1'].fill = GREEN
    ws['A2'] = L['subs']; ws['A3'] = 'Goal: ' + L['goal']; ws['A3'].font = Font(bold=True)
    ws['A4'] = 'Excel tools: ' + ', '.join(L['fx'])
    for i, s in enumerate(L['steps']):
        ws.cell(row=6 + i, column=1, value=f'{i+1}. ' + strip(s)).alignment = Alignment(wrap_text=True)
    r = 7 + len(L['steps'])
    ws.cell(row=r, column=1, value='Type your answers in the yellow cells on the Tasks sheet. Save the file as Lab' + L['week'].replace(' ', '') + '_YourName.xlsx and upload it on Moodle.').font = Font(italic=True)
    ws.cell(row=r + 1, column=1, value='Data are illustrative. Instructor: Eng. Nadia Mek.').font = Font(italic=True, color='777777')
    ws.column_dimensions['A'].width = 110
    if L.get('data'):
        h, rows = L['data']; d = wb.create_sheet('Data'); d.append(h)
        for c in d[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = GREEN
        for row in rows: d.append(row)
        d.freeze_panes = 'A2'; autosize(d)
    if L.get('sim'):
        s = wb.create_sheet('Sim'); s['A1'] = 'Build the simulation here (see Instructions).'; s['A1'].font = Font(italic=True)
    if L.get('variables'):
        v = wb.create_sheet('Variables'); v.append(['Variable', 'Type (choose)', 'Example value'])
        for c in v[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = GREEN
        from openpyxl.worksheet.datavalidation import DataValidation
        dv = DataValidation(type='list', formula1='"Nominal,Ordinal,Discrete,Continuous"', allow_blank=True); v.add_data_validation(dv)
        for i, name in enumerate(SH):
            v.append([name, None, SROWS[0][i]]); dv.add(f'B{i+2}'); v[f'B{i+2}'].fill = YEL
        autosize(v)
    t = wb.create_sheet('Tasks'); t.append(['Task', 'Your answer']);
    for c in t[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = GREEN
    for q in L['tasks']:
        t.append([q, None]); t.cell(row=t.max_row, column=2).fill = YEL
    t.column_dimensions['A'].width = 70; t.column_dimensions['B'].width = 30
    wb.save(os.path.join(ROOT, 'lab', L['id'] + '.xlsx'))
    if L.get('csv'):
        import csv
        h, rows = L['data']
        with open(os.path.join(ROOT, 'lab', 'lab-01-students.csv'), 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f); w.writerow(h); w.writerows(rows)
    print('wrote', L['id'])

web = [{k: L[k] for k in ('id', 'week', 'title', 'subs', 'clo', 'assess', 'goal', 'fx', 'steps', 'grid')} | {'csv': bool(L.get('csv'))} for L in LABS]
open(os.path.join(ROOT, 'assets', 'lab-data.js'), 'w', encoding='utf-8').write('/* generated by tools/lab.py */\nwindow.MTH105_LABS=' + json.dumps(web, ensure_ascii=False) + ';\n')
print('lab-data.js', len(web))
